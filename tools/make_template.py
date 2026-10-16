from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
import sys
COLS=['الدرجة','الجنس','الطالب','العنوان','اليوم','التاريخ','الوقت','المكان','ملاحظة المكان','رابط البث']
for i in range(1,6): COLS+= [f'عضو{i} الاسم',f'عضو{i} الصفة',f'عضو{i} الدور']
NOTES={'الكلية':'مثال: يسر كلية الشريعة','نص الإعلان':'مثال: الإعلان عن مناقشة رسالة علمية','الدرجة':'مثال: الماجستير في الفقه المقارن',
 'الجنس':'اختر: ذكر أو أنثى (يغيّر «يقدمها/تقدمها» و«طالب/طالبة»)','الطالب':'اسم الطالب أو الطالبة',
 'العنوان':'عنوان الرسالة. لسطر جديد اكتب | مثال: العنوان | (دراسة تطبيقية)','اليوم':'مثال: الأحد','التاريخ':'مثال: 1448/4/16 هـ',
 'الوقت':'مثال: 8 صباحاً','المكان':'مثال: قاعة ابن عثيمين','ملاحظة المكان':'مثال: (والبث عبر الإنترنت)',
 'رابط البث':'الافتراضي: قناة اليوتيوب https://www.youtube.com/@src.sds1 (اتركه فاضي أو اختره من القائمة). تقدر تلصق رابط زوم أو بلاك بورد أو أي منصة.'}
teal='2B7574'; dark='0F4A4D'; light='D9EAE9'; mem='E6EFE0'; memh='4F7F5E'
thin=Side(style='thin',color='C5DBDA'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); ws=wb.active; ws.title='الإعلانات'; ws.sheet_view.rightToLeft=True
ws.append(COLS)
for j,c in enumerate(COLS,1):
    cell=ws.cell(1,j); m=c.startswith('عضو')
    cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=11)
    cell.fill=PatternFill('solid',fgColor=memh if m else teal)
    cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    cell.border=border
    note = NOTES.get(c) or ('اسم العضو مع اللقب، مثال: أ.د. عمر بن علي السديس' if c.endswith('الاسم') else 'مثال: الأستاذ في قسم الفقه' if c.endswith('الصفة') else 'مثال: مشرفاً ومقرراً / مناقشاً')
    cell.comment=Comment(note,'النموذج'); cell.comment.width=260; cell.comment.height=60
    w = 55 if c=='العنوان' else 34 if c=='رابط البث' else 12 if c in ('الجنس','اليوم','الوقت') else 24
    ws.column_dimensions[L(j)].width=w
ws.row_dimensions[1].height=32
N=300
for r in range(2,N+2):
    for j,c in enumerate(COLS,1):
        cell=ws.cell(r,j); cell.border=border; cell.number_format='@'
        cell.font=Font(name='Arial',size=11); cell.alignment=Alignment(vertical='center',wrap_text=(c=='العنوان'))
        if c.startswith('عضو') and r%2==0: cell.fill=PatternFill('solid',fgColor='F4F8F1')
        elif r%2==0: cell.fill=PatternFill('solid',fgColor='F3F8F8')
ws.freeze_panes='F2'
def dv(formula, col, prompt):
    v=DataValidation(type='list',formula1=formula,allow_blank=True,showErrorMessage=False,showInputMessage=True,promptTitle='اختيار',prompt=prompt)
    ws.add_data_validation(v); v.add(f'{col}2:{col}{N+1}')
idx=lambda c: L(COLS.index(c)+1)
dv('"ذكر,أنثى"', idx('الجنس'), 'ذكر أو أنثى')
dv('"https://www.youtube.com/@src.sds1"', idx('رابط البث'), 'الافتراضي قناة اليوتيوب — أو الصق أي رابط آخر')
dv('"الماجستير في الفقه المقارن,الماجستير في الفقه,الماجستير في أصول الفقه,الدكتوراه في الفقه,الدكتوراه في الفقه المقارن,الدكتوراه في أصول الفقه"', idx('الدرجة'), 'اختر أو اكتب الدرجة')
dv('"قاعة ابن عثيمين,القاعة الكبرى بمبنى الطالبات"', idx('المكان'), 'اختر أو اكتب المكان')
dv('"الأحد,الاثنين,الثلاثاء,الأربعاء,الخميس,الجمعة,السبت"', idx('اليوم'), 'اختر اليوم')
for i in range(1,6):
    dv('"مشرفاً ومقرراً,مشرفاً,مشرفاً مشاركاً,مناقشاً,مناقشاً خارجياً,مشرفةً ومقررةً,مناقِشةً"', idx(f'عضو{i} الدور'), 'اختر من القائمة أو اكتب دوراً آخر')
dv('"الأستاذ في قسم الفقه,الأستاذ المشارك في قسم الفقه,الأستاذ المساعد في قسم الفقه"', idx('عضو1 الصفة'), 'اختر أو اكتب الصفة')

# ورقة الشرح
h=wb.create_sheet('طريقة التعبئة'); h.sheet_view.rightToLeft=True
h.column_dimensions['A'].width=110
lines=[('طريقة تعبئة النموذج','t'),
('1) عبّئ ورقة «الإعلانات»: كل سطر = إعلان مستقل. لا تغيّر أسماء الأعمدة في السطر الأول.',''),
('2) «يسر كلية الشريعة» ونص الإعلان ثابتة في التصميم. أي عمود تتركه فاضي يأخذ قيمته من التصميم الحالي في الموقع (مثل المكان).',''),
('3) الجنس: اختر «ذكر» أو «أنثى» — يغيّر «يقدمها/تقدمها» و«طالب/طالبة الدراسات العليا» تلقائياً.',''),
('4) العنوان: لبدء سطر جديد داخل العنوان اكتب العلامة |   مثال: المسائل الفقهية ... الحنبلي | (جمعاً ودراسة)',''),
('5) لجنة المناقشة: عضو1 ثم عضو2 ... حتى عضو5 (الاسم، الصفة، الدور). اترك الأعضاء غير المطلوبين فاضيين.',''),
('6) رابط البث: الافتراضي قناة اليوتيوب (اختره من القائمة أو اتركه فاضي). تقدر تلصق رابط زوم أو بلاك بورد أو أي منصة. يظهر كرمز QR إذا فعّلته في الموقع.',''),
('7) في الموقع: «جدول بيانات الإعلانات» ← «📄 من ملف Excel / CSV» ← اختر هذا الملف. ستظهر البيانات في الجدول وتقدر تعدّلها ثم «📦 تحميل المحدد».',''),
('',''),
('استخدامه في Google Sheets','t'),
('ارفع هذا الملف إلى Google Drive ثم افتحه بـ Google Sheets (أو: ملف ← استيراد). بعد التعبئة: «مشاركة» ← «أي شخص لديه الرابط» (عارض)،',''),
('ثم انسخ الرابط والصقه في الموقع واضغط «🔗 من Google Sheets». وكل ما عدّلت الشيت اضغط «🔄 تحديث» في الموقع.',''),
('',''),('مثال على سطر معبّأ','t')]
for i,(t,k) in enumerate(lines,1):
    c=h.cell(i,1,t); c.font=Font(name='Arial',size=13 if k else 11,bold=bool(k),color=dark if k else '222222'); c.alignment=Alignment(wrap_text=True,vertical='center')
    h.row_dimensions[i].height=26 if k else 22
ex=['الماجستير في الفقه المقارن','ذكر','عبد الله بن عادل السعدي',
 'المسائل الفقهية التي خالف فيها الإمام محمد بن عبد الوهاب المذهب الحنبلي | (جمعاً ودراسة)','الأحد','1448/4/16 هـ','8 صباحاً','قاعة ابن عثيمين','(والبث عبر الإنترنت)','https://www.youtube.com/@src.sds1',
 'أ.د. عمر بن علي السديس','الأستاذ في قسم الفقه','مشرفاً ومقرراً','أ.د. فهد بن صالح الحمود','الأستاذ في قسم الفقه','مناقشاً','د. أحمد بن سليمان العودة','الأستاذ المساعد في قسم الفقه','مناقشاً']
r0=len(lines)+1
for j,(c,v) in enumerate(zip(COLS,ex)):
    a=h.cell(r0+j,1,f'{c}:   {v}'); a.font=Font(name='Arial',size=11); a.fill=PatternFill('solid',fgColor='F3F8F8' if j%2 else 'FFFFFF')
wb.active=0
wb.save(sys.argv[1]); print('saved')
